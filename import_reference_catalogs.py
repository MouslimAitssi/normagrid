"""Charge les catalogues Excel de charges et de cables dans le projet SQLite actif."""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from backend.db import get_connection


def database_value(value):
    """Valeur compatible SQLite et JSON, y compris les booleens Excel."""
    if isinstance(value, bool):
        return int(value)
    return value


def visible_number(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_rows(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    return [
        {headers[index]: database_value(value) for index, value in enumerate(row)}
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in row)
    ]


def import_charges(conn, path):
    rows = read_rows(path)
    conn.execute("DELETE FROM reference_charges")
    values = []
    for index, row in enumerate(rows, start=1):
        designation = (
            f"{row['TypeConso']} — {row['Consom']} "
            f"({visible_number(row['UnMin'])}-{visible_number(row['UnMax'])} V)"
        )
        values.append((
            index, designation, row["TypeConso"], row["Consom"], row["TypRecept"],
            row["PElectrique"], row["Polarite"], row["Rendmnt"], row["INominal"],
            row["Is50Hz"], row["Is60Hz"], row["IsDC"], row["KUtilisation"],
            row["KFoison"], row["CosPhi"], row["CosPhiDem"], row["IdSurIn"],
            row["UnMin"], row["UnMax"],
        ))
    conn.executemany(
        """
        INSERT INTO reference_charges (
            id, designation, type_conso, consom, typ_recept, p_electrique_w,
            polarite, rendement, i_nominal, is_50hz, is_60hz, is_dc,
            k_utilisation, k_foison, cos_phi, cos_phi_dem, id_sur_in, un_min, un_max
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        values,
    )
    return len(values)


def import_cables(conn, path):
    rows = read_rows(path)
    conn.execute("DELETE FROM reference_cables")
    values = []
    for row in rows:
        designation = f"{row['FamilleCable']} {row['Section']}"
        values.append((
            row["ID"], designation, row["FamilleCable"], row["Section"], row["SectionTxt"],
            row["SectionReelle"], row["NbConducteur"], row["MetalAme"], row["IsArme"],
            row["TypeConduct"], row["VertJaune"], row["FamilleCuAl"], row["TempMax"],
            row["Diametre"], row["Poids"], row["IzAir"], row["Un"], row["UnMax"],
            json.dumps(row, ensure_ascii=False),
        ))
    conn.executemany(
        """
        INSERT INTO reference_cables (
            id, designation, famille_cable, section, section_txt, section_reelle,
            nb_conducteur, metal_ame, is_arme, type_conduct, vert_jaune,
            famille_cu_al, temp_max, diametre, poids, iz_air, un, un_max, raw_data_json
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        values,
    )
    return len(values)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--charges", required=True, type=Path, help="Chemin vers Table Charge.xlsx")
    parser.add_argument("--cables", required=True, type=Path, help="Chemin vers Table Cable.xlsx")
    args = parser.parse_args()

    conn = get_connection()
    try:
        charge_count = import_charges(conn, args.charges)
        cable_count = import_cables(conn, args.cables)
        conn.commit()
    finally:
        conn.close()
    print(f"Catalogues importes : {charge_count} charges, {cable_count} cables.")


if __name__ == "__main__":
    main()
