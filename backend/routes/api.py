from flask import Blueprint, jsonify, request
from backend import models
import re

api = Blueprint("api", __name__, url_prefix="/api")


@api.get("/schema")
def schema():
    return jsonify(models.get_schema())


@api.get("/tree")
def tree():
    return jsonify(models.get_tree())


@api.get("/graph")
def graph():
    return jsonify(models.get_graph())


@api.get("/cable-schedule")
def cable_schedule():
    return jsonify(models.get_cable_schedule())


@api.get("/bilan-puissance")
def bilan_puissance():
    return jsonify(models.get_bilan_puissance())


@api.post("/bilan-puissance/export")
def bilan_puissance_export():
    import os
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    data = request.get_json(force=True) or {}
    folder = (data.get("folder") or "").strip()
    filename = (data.get("filename") or "").strip()
    if not folder:
        return jsonify({"error": "Le dossier de destination est obligatoire."}), 400
    if not filename:
        return jsonify({"error": "Le nom du fichier est obligatoire."}), 400
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"

    try:
        os.makedirs(folder, exist_ok=True)
    except OSError as e:
        return jsonify({"error": f"Impossible de creer/acceder au dossier : {e}"}), 400

    bilan = models.get_bilan_puissance()
    all_rows = bilan["rows"]
    tension_ht_kv = bilan["tension_ht_kv"]

    groups = {}
    for r in all_rows:
        key = r.get("switchgear_mcc") or "Non affecte"
        groups.setdefault(key, []).append(r)

    headers = [
        "Designation", "Switchgear / MCC", "Equipment Tag No", "E/NE", "% E", "Type",
        "Service I/C", "W/S", "P. Apparente Unit. (kVA)", "P. Installee", "Unite", "Rendement",
        "P. Active Unit. (kW)", "Nombre", "Cos φ", "P. Reactive Unit. (kvar)", "K.Util", "K.Simul",
        "P. Active Foisonnee (kW)", "P. Reactive Foisonnee (kvar)", "Alimentation (V)",
    ]
    field_order = [
        "designation", "switchgear_mcc", "equipment_tag_no", "e_ne", "pct_e", "type_equipement",
        "c_i", "w_s", "puissance_apparente_unitaire", "puissance_installee", "unite", "rendement",
        "puissance_active_unitaire", "nombre", "cos_phi", "puissance_reactive_unitaire", "k_util",
        "k_simul", "puissance_active_foisonnee", "puissance_reactive_foisonnee", "alimentation",
    ]

    header_fill = PatternFill("solid", start_color="1F3864", end_color="1F3864")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri")
    body_font = Font(name="Calibri", size=10)
    kpi_fill = PatternFill("solid", start_color="EAF2F8", end_color="EAF2F8")
    kpi_font = Font(name="Calibri", size=10, bold=True)

    def write_sheet(ws, rows_for_sheet):
        kpis = models._compute_bilan_kpis(rows_for_sheet, tension_ht_kv)
        kpi_labels = [
            ("kW installes total", kpis["kw_installes"]),
            ("kW absorbes total", kpis["kw_absorbes"]),
            ("kVA reseau total", kpis["kva_reseau"]),
            ("kVAr total", kpis["kvar_total"]),
            ("cos \u03c6 global", kpis["cos_phi_global"]),
            ("I reseau HTA (A)", kpis["i_reseau_hta"]),
        ]
        for i, (label, value) in enumerate(kpi_labels):
            ws.cell(row=1, column=2 * i + 1, value=label).font = kpi_font
            ws.cell(row=1, column=2 * i + 1).fill = kpi_fill
            ws.cell(row=2, column=2 * i + 1, value=value if value is not None else "")
            ws.cell(row=1, column=2 * i + 1).alignment = Alignment(horizontal="center")
            ws.cell(row=2, column=2 * i + 1).alignment = Alignment(horizontal="center")

        header_row = 4
        for j, h in enumerate(headers):
            cell = ws.cell(row=header_row, column=j + 1, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        for r_idx, row in enumerate(rows_for_sheet, start=header_row + 1):
            for j, field in enumerate(field_order):
                v = row.get(field)
                cell = ws.cell(row=r_idx, column=j + 1, value=v if v is not None else "")
                cell.font = body_font
        ws.freeze_panes = f"A{header_row + 1}"
        for j in range(len(headers)):
            col_letter = get_column_letter(j + 1)
            length = max(
                [len(str(headers[j]))] +
                [len(str(row.get(field_order[j]) or "")) for row in rows_for_sheet]
            )
            ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 32)

    wb = Workbook()
    sheet_names_used = set()

    def unique_sheet_name(name):
        base = re.sub(r"[\[\]\*/\\\?:]", "_", str(name))[:28] or "Feuille"
        candidate = base
        i = 2
        while candidate in sheet_names_used:
            candidate = f"{base[:25]}_{i}"
            i += 1
        sheet_names_used.add(candidate)
        return candidate

    first = True
    for key in sorted(groups.keys()):
        ws = wb.active if first else wb.create_sheet()
        ws.title = unique_sheet_name(key)
        first = False
        write_sheet(ws, groups[key])

    if not groups:
        wb.active.title = "Bilan"

    full_path = os.path.join(folder, filename)
    try:
        wb.save(full_path)
    except OSError as e:
        return jsonify({"error": f"Impossible d'ecrire le fichier : {e}"}), 400

    return jsonify({"ok": True, "path": full_path, "nb_sheets": len(groups) or 1}), 201


@api.post("/cable-schedule/export")
def cable_schedule_export():
    import os
    import re
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    data = request.get_json(force=True) or {}
    folder = (data.get("folder") or "").strip()
    filename = (data.get("filename") or "").strip()

    if not folder:
        return jsonify({"error": "Le dossier de destination est obligatoire."}), 400
    if not filename:
        return jsonify({"error": "Le nom du fichier est obligatoire."}), 400
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"

    try:
        os.makedirs(folder, exist_ok=True)
    except OSError as e:
        return jsonify({"error": f"Impossible de creer/acceder au dossier : {e}"}), 400

    full_path = os.path.join(folder, filename)

    rows = models.get_cable_schedule()

    wb = Workbook()
    ws = wb.active
    ws.title = "Carnet de cables"
    headers = [
        "Tenant (origine)", "Type origine", "Aboutissant (destination)", "Type destination",
        "Designation", "Type de cable", "Longueur (m)", "Protection", "Calibre (A)", "Differentiel (mA)",
    ]
    ws.append(headers)
    for r in rows:
        ws.append([
            r.get("tenant") or "", r.get("type_origine") or "",
            r.get("aboutissant") or "", r.get("type_destination") or "",
            r.get("designation") or "", r.get("type_cable") or "",
            r.get("longueur_m") if r.get("longueur_m") is not None else "",
            r.get("protection_modele") or "",
            r.get("calibre_a") if r.get("calibre_a") is not None else "",
            r.get("differentiel_ma") if r.get("differentiel_ma") is not None else "",
        ])

    header_fill = PatternFill("solid", start_color="1F3864", end_color="1F3864")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri")
    body_font = Font(name="Calibri", size=10)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 45)

    try:
        wb.save(full_path)
    except OSError as e:
        return jsonify({"error": f"Impossible d'ecrire le fichier : {e}"}), 400

    return jsonify({"ok": True, "path": full_path, "nb_rows": len(rows)}), 201


@api.get("/<table>")
def list_table(table):
    if table not in models.TABLES:
        return jsonify({"error": f"table inconnue: {table}"}), 404
    return jsonify(models.list_rows(table))


@api.post("/<table>")
def create_row(table):
    if table not in models.TABLES:
        return jsonify({"error": f"table inconnue: {table}"}), 404
    data = request.get_json(force=True)
    try:
        models.insert_row(table, data)
    except Exception as e:
        return jsonify({"error": str(e)}), 400
    return jsonify({"ok": True}), 201


@api.get("/<table>/<path:pk>")
def get_row(table, pk):
    if table not in models.TABLES:
        return jsonify({"error": f"table inconnue: {table}"}), 404
    pk_values = pk.split("|")
    row = models.get_row(table, pk_values)
    if row is None:
        return jsonify({"error": "entree introuvable"}), 404
    return jsonify(row)


@api.put("/<table>/<path:pk>")
def update_row(table, pk):
    if table not in models.TABLES:
        return jsonify({"error": f"table inconnue: {table}"}), 404
    pk_values = pk.split("|")
    data = request.get_json(force=True)
    try:
        models.update_row(table, pk_values, data)
    except Exception as e:
        return jsonify({"error": str(e)}), 400
    return jsonify({"ok": True})


@api.delete("/<table>/<path:pk>")
def delete_row(table, pk):
    if table not in models.TABLES:
        return jsonify({"error": f"table inconnue: {table}"}), 404
    pk_values = pk.split("|")
    try:
        models.delete_row(table, pk_values)
    except Exception as e:
        return jsonify({"error": str(e)}), 400
    return jsonify({"ok": True})
